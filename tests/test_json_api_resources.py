from io import BytesIO

from openpyxl import load_workbook

from app.models import Faculty, Specialty


def test_json_api_core_resources_are_executable(client, session, organization):
    faculty = Faculty(name="Автотракторный", code="АТФ")
    specialty = Specialty(
        code="1-37 01 03",
        name="Тракторостроение",
        profile="Проектирование тракторов",
        qualification="Инженер-механик",
    )
    session.add_all([faculty, specialty])
    session.commit()

    created_contract = client.post("/api/contracts", json={
        "organization_id": organization.id,
        "number": "221-АТФ/280",
        "date_start": "2020-10-01",
        "date_end": "2030-12-31",
        "faculty_ids": [faculty.id],
    })
    assert created_contract.status_code == 201, created_contract.text
    contract_id = created_contract.json()["id"]
    assert created_contract.json()["current_order"] is None

    item = client.post(f"/api/contracts/{contract_id}/order-items", json={
        "faculty_id": faculty.id,
        "specialty_id": specialty.id,
        "profile": "Тракторы",
        "qualification": "Инженер-механик",
        "years": {"2026": 15, "2027": 16},
    })
    assert item.status_code == 201, item.text
    assert item.json()["years"] == {"2026": 15, "2027": 16}

    contract = client.get(f"/api/contracts/{contract_id}")
    assert contract.status_code == 200
    assert contract.json()["number"] == "221-АТФ/280"
    assert contract.json()["current_order"]["items"][0]["years"]["2026"] == 15

    created_application = client.post("/api/applications", json={
        "organization_id": organization.id,
        "number": "З-2026/15",
        "signed_date": "2026-09-25",
        "date_end": "2027-09-25",
        "faculty_ids": [faculty.id],
    })
    assert created_application.status_code == 201, created_application.text
    application_id = created_application.json()["id"]
    assert created_application.json()["status"] == "Заявка"

    organization_card = client.get(f"/api/organizations/{organization.id}")
    assert organization_card.status_code == 200
    assert {row["id"] for row in organization_card.json()["contracts"]} == {contract_id}
    assert {row["id"] for row in organization_card.json()["applications"]} == {application_id}

    contracts = client.get("/api/contracts", params=[("faculty_ids", faculty.id)])
    applications = client.get("/api/applications", params=[("faculty_ids", faculty.id)])
    assert contracts.status_code == applications.status_code == 200
    assert contracts.json()["total"] == applications.json()["total"] == 1

    specialty_update = client.put(f"/api/specialties/{specialty.id}", json={
        "name": "Тракторостроение",
        "profile": "",
        "qualification": "Инженер",
    })
    assert specialty_update.status_code == 200
    assert specialty_update.json()["profile"] is None

    metadata = client.get("/api/audit/metadata")
    assert metadata.status_code == 200
    assert metadata.json()["actions"] and metadata.json()["entity_types"]

    exported = client.get("/api/export/contracts", params=[("faculty_ids", faculty.id)])
    assert exported.status_code == 200
    assert load_workbook(BytesIO(exported.content)).active.max_row == 2


def test_json_admin_user_endpoints_and_validation(client):
    created = client.post("/api/users", json={
        "full_name": "Петров Пётр Петрович",
        "username": "petrov",
        "email": "petrov@bntu.by",
        "role": "HEAD",
        "initial_password": "temporary-2026",
    })
    assert created.status_code == 201, created.text
    user_id = created.json()["id"]
    assert created.json()["email"] == "petrov@bntu.by"
    assert created.json()["need_password_change"] is True

    duplicate = client.post("/api/users", json={
        "full_name": "Другой сотрудник",
        "username": "another",
        "email": "petrov@bntu.by",
        "role": "HEAD",
        "initial_password": "temporary-2026",
    })
    assert duplicate.status_code == 409
    assert "электронной почтой" in duplicate.json()["detail"]

    reset = client.post(f"/api/users/{user_id}/reset-password", json={"new_password": "new-password-2026"})
    assert reset.status_code == 204
    assert reset.content == b""


def test_json_validation_errors_are_readable(client):
    invalid = client.post("/api/contracts", json={
        "organization_id": 999,
        "number": "",
        "date_start": "not-a-date",
        "faculty_ids": [],
    })
    assert invalid.status_code == 422
    assert set(invalid.json()) == {"detail"}
    assert invalid.json()["detail"]


def test_contract_order_item_without_qualification_is_returned_as_null(client, session, contract):
    specialty = Specialty(code="6-05-0718-01", name="6-05-0718-01", qualification=None)
    session.add(specialty)
    session.commit()
    faculty_id = contract.faculty_links[0].faculty_id

    created = client.post(f"/api/contracts/{contract.id}/order-items", json={
        "faculty_id": faculty_id,
        "specialty_id": specialty.id,
        "profile": None,
        "qualification": None,
        "years": {"2026": 4},
    })
    assert created.status_code == 201, created.text
    assert created.json()["qualification"] is None

    card = client.get(f"/api/contracts/{contract.id}")
    assert card.status_code == 200
    assert card.json()["current_order"]["items"][0]["qualification"] is None
